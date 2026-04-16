"""
Fill the data-inventarisatie xlsx using the same column layout as
the user's template (`data-inventarisatie-ingevuld.xlsx`).

Columns:
    Processtap | Stap-ID | Dataobject | Attribuut | Verplicht? |
    Doelbinding | Classificatie | Autorisatie (wie?) | Bewaartermijn |
    Bron (master) | Opmerkingen
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from merger import MergedModel

HEADERS = [
    "Processtap", "Stap-ID", "Dataobject", "Attribuut", "Verplicht?",
    "Doelbinding", "Classificatie", "Autorisatie (wie?)", "Bewaartermijn",
    "Bron (master)", "Opmerkingen",
]

CLASSIFICATION_FILL = {
    "Bijzonder persoonsgegeven": PatternFill("solid", fgColor="FFEACE"),
    "Vertrouwelijk":             PatternFill("solid", fgColor="FFF2CC"),
    "Intern":                    PatternFill("solid", fgColor="E1F0E1"),
    "Openbaar":                  PatternFill("solid", fgColor="DDEBF7"),
}

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(ws, row: int = 1) -> None:
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _autosize(ws) -> None:
    widths = [28, 9, 22, 26, 11, 32, 22, 22, 22, 18, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_row(ws, row: int, r) -> None:
    values = [
        r.process_step, r.step_id, r.data_object, r.attribute, r.required,
        r.purpose, r.classification, r.authorization, r.retention,
        r.source, r.remarks,
    ]
    fill = CLASSIFICATION_FILL.get(r.classification)
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            c.fill = fill


def write_xlsx(model: MergedModel, out_path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    # --- Sheet per BPMN ---
    rows_by_process: dict[str, list] = {}
    for r in model.inventory:
        rows_by_process.setdefault(r.process, []).append(r)

    # Combined sheet first
    ws_all = wb.create_sheet("Alle processen")
    _write_header(ws_all)
    _autosize(ws_all)
    row_idx = 2
    for process, rows in rows_by_process.items():
        # Section header per process
        c = ws_all.cell(row=row_idx, column=1,
                        value=f"=== {process} ({len(rows)} regels) ===")
        c.font = Font(bold=True, size=11)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        ws_all.merge_cells(start_row=row_idx, start_column=1,
                           end_row=row_idx, end_column=len(HEADERS))
        row_idx += 1
        for r in rows:
            _write_row(ws_all, row_idx, r)
            row_idx += 1
    ws_all.freeze_panes = "A2"

    # Per-process sheet
    for process, rows in rows_by_process.items():
        # sheet names limited to 31 chars
        name = (process[:28] + "…") if len(process) > 31 else process
        ws = wb.create_sheet(name)
        _write_header(ws)
        _autosize(ws)
        for i, r in enumerate(rows, start=2):
            _write_row(ws, i, r)
        ws.freeze_panes = "A2"

    # --- Anker-objecten sheet ---
    ws_anchor = wb.create_sheet("Ankerobjecten")
    headers = ["Ankerobject", "Aantal BPMN's", "Voorkomt in"]
    for col, h in enumerate(headers, start=1):
        c = ws_anchor.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    ws_anchor.column_dimensions["A"].width = 36
    ws_anchor.column_dimensions["B"].width = 14
    ws_anchor.column_dimensions["C"].width = 80
    for i, name in enumerate(model.anchor_objects(), start=2):
        files = sorted({sf for sf, _ in model.data_object_index[name.lower()]})
        ws_anchor.cell(row=i, column=1, value=name).font = Font(bold=True)
        ws_anchor.cell(row=i, column=2, value=len(files))
        ws_anchor.cell(row=i, column=3, value=", ".join(files)).alignment = (
            Alignment(wrap_text=True, vertical="top")
        )

    # --- Actoren sheet ---
    ws_act = wb.create_sheet("Actoren")
    for col, h in enumerate(["Actor", "Type", "Voorkomt in", "Onderbouwing"],
                            start=1):
        c = ws_act.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    ws_act.column_dimensions["A"].width = 38
    ws_act.column_dimensions["B"].width = 14
    ws_act.column_dimensions["C"].width = 60
    ws_act.column_dimensions["D"].width = 60
    for i, actor in enumerate(model.actors, start=2):
        ws_act.cell(row=i, column=1, value=actor.name).font = Font(bold=True)
        ws_act.cell(row=i, column=2, value=actor.subtype)
        ws_act.cell(row=i, column=3,
                    value=", ".join(actor.evidence.get("appears_in", [])))
        ws_act.cell(row=i, column=4,
                    value=actor.evidence.get("classification_reason", "")
                    ).alignment = Alignment(wrap_text=True, vertical="top")

    # --- Legenda ---
    ws_leg = wb.create_sheet("Legenda")
    legenda = [
        ("Classificatieniveaus", ""),
        ("Openbaar", "Iedereen mag dit zien"),
        ("Intern", "Alleen FNV-medewerkers"),
        ("Vertrouwelijk", "Specifieke rollen (bijv. IBAN, salaris)"),
        ("Bijzonder persoonsgegeven",
         "Vakbondslidmaatschap, gezondheid, etnische afkomst (AVG art. 9)"),
        ("", ""),
        ("Methodiek extractie", ""),
        ("Processtap", "Afgeleid uit <bpmn:task> en alle subtypes "
                       "(userTask, serviceTask, manualTask, ...)."),
        ("Dataobject", "Afgeleid uit <bpmn:dataObject> en "
                       "<bpmn:dataObjectReference>."),
        ("Bron (master)", "Default 'BPMN dataObject'; bij explicite "
                          "<bpmn:dataStore(Reference)> wordt dit overschreven."),
        ("Procesattribuut", "Afgeleid uit gateways en condition expressions."),
        ("Ankerobject", "Data-object met dezelfde naam in ≥2 BPMN-bestanden."),
        ("Actor (intern)", "Afgeleid uit <bpmn:lane>."),
        ("Actor (extern)", "Afgeleid uit <bpmn:participant> zonder processRef."),
    ]
    for i, (k, v) in enumerate(legenda, start=1):
        ws_leg.cell(row=i, column=1, value=k).font = Font(
            bold=k.endswith(":") or k in ("Classificatieniveaus",
                                          "Methodiek extractie"))
        ws_leg.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True)
    ws_leg.column_dimensions["A"].width = 28
    ws_leg.column_dimensions["B"].width = 70

    wb.save(out_path)
